"""End-to-end integration test: Audio Slicing -> Whisper -> Word .docx export."""

import subprocess
import time
from pathlib import Path
import pytest
from docx import Document

from whisper_gui.audio_utils import get_audio_info
from whisper_gui.models import AudioTrack, SliceConfig, SliceMode, TranscriptionConfig
from whisper_gui.transcriber import Transcriber
from whisper_gui.worker import TranscriptionWorker


@pytest.fixture
def test_audio_file(tmp_path: Path) -> Path:
    """Create a 3-second test WAV file with sine wave."""
    wav_file = tmp_path / "test_audio.wav"
    cmd = [
        "ffmpeg", "-y",
        "-f", "lavfi",
        "-i", "sine=frequency=440:duration=3",
        str(wav_file),
    ]
    subprocess.run(cmd, check=True, capture_output=True)
    return wav_file


def test_worker_end_to_end(qtbot, test_audio_file: Path, tmp_path: Path):
    info = get_audio_info(test_audio_file)
    track = AudioTrack(test_audio_file, info)

    # Slice 1 second (from 0.5s to 1.5s)
    track.slice_config = SliceConfig(
        mode=SliceMode.RANGE,
        start_sec=0.5,
        end_sec=1.5,
    )

    out_docx = tmp_path / "output_test.docx"
    track.output_docx_path = out_docx

    # Configure transcription with 'tiny' for fast test execution
    cfg = TranscriptionConfig(
        model_name="tiny",
        device="cpu",
        compute_type="int8",
        cpu_threads=2,
        vad_filter=False,  # Sine tone has no voice, vad would drop everything
    )

    transcriber = Transcriber()
    worker = TranscriptionWorker(
        tracks=[track],
        transcription_config=cfg,
        transcriber=transcriber,
    )

    completed_tracks = []
    def on_completed(track_id, result):
        completed_tracks.append((track_id, result))

    worker.track_completed.connect(on_completed)

    with qtbot.waitSignal(worker.batch_finished, timeout=15000):
        worker.start()

    assert len(completed_tracks) == 1
    track_id, result = completed_tracks[0]
    assert track_id == track.id
    assert out_docx.exists()

    # Verify Word document
    doc = Document(str(out_docx))
    assert len(doc.tables) >= 1
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Audio Transcription" in doc_text
    assert "test_audio.wav" in "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)


def test_worker_with_llm_polish(qtbot, test_audio_file: Path, tmp_path: Path):
    import json
    from unittest.mock import MagicMock, patch
    from whisper_gui.models import LLMConfig

    info = get_audio_info(test_audio_file)
    track = AudioTrack(test_audio_file, info)
    out_docx = tmp_path / "output_polished.docx"
    track.output_docx_path = out_docx

    llm_cfg = LLMConfig(
        enabled=True,
        provider="DeepSeek",
        model_name="deepseek-v4.1-flash",
        api_key="sk-test",
    )
    cfg = TranscriptionConfig(
        model_name="tiny",
        device="cpu",
        compute_type="int8",
        cpu_threads=2,
        vad_filter=False,
        llm_config=llm_cfg,
    )

    mock_resp = MagicMock()
    mock_resp.read.return_value = json.dumps({
        "choices": [{"message": {"role": "assistant", "content": "Polished audio transcript from DeepSeek."}}]
    }).encode("utf-8")

    from whisper_gui.models import TranscriptSegment, TranscriptionResult
    mock_whisper_result = TranscriptionResult(
        segments=[TranscriptSegment(start=0.0, end=2.0, text="Raw speech from audio")],
        full_text="Raw speech from audio",
        audio_duration=2.0,
        elapsed_time=1.0,
    )

    transcriber = Transcriber()
    worker = TranscriptionWorker(
        tracks=[track],
        transcription_config=cfg,
        transcriber=transcriber,
    )

    completed_tracks = []
    worker.track_completed.connect(lambda t_id, res: completed_tracks.append((t_id, res)))

    with patch.object(transcriber, "transcribe", return_value=mock_whisper_result), patch("urllib.request.urlopen") as mock_url:
        mock_url.return_value.__enter__.return_value = mock_resp
        with qtbot.waitSignal(worker.batch_finished, timeout=15000):
            worker.start()

    assert len(completed_tracks) == 1
    _, result = completed_tracks[0]
    assert result.polished_text == "Polished audio transcript from DeepSeek."
    assert result.llm_model_used == "deepseek-v4.1-flash"
    assert out_docx.exists()

    doc = Document(str(out_docx))
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Polished & Refined Document" in doc_text
    assert "Polished audio transcript from DeepSeek." in doc_text
