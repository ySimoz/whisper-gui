"""Unit tests for Word (.docx) export generation."""

from pathlib import Path
from docx import Document

from whisper_gui.docx_exporter import export_transcript_to_docx
from whisper_gui.models import (
    AudioInfo,
    SliceConfig,
    SliceMode,
    TranscriptSegment,
    TranscriptionConfig,
    TranscriptionResult,
)


def test_export_transcript_to_docx_with_timestamps(tmp_path: Path):
    docx_file = tmp_path / "transcript_test.docx"

    track_info = AudioInfo(
        path=Path("/music/lecture.mp3"),
        filename="lecture.mp3",
        duration=1800.0,
        format_name="mp3",
        sample_rate=44100,
        channels=2,
    )
    slice_cfg = SliceConfig(mode=SliceMode.RANGE, start_sec=60.0, end_sec=360.0)
    tx_cfg = TranscriptionConfig(
        model_name="large-v3-turbo",
        language="it",
        include_timestamps_in_docx=True,
    )

    segments = [
        TranscriptSegment(start=0.0, end=4.5, text="Buongiorno a tutti."),
        TranscriptSegment(start=5.0, end=10.2, text="Oggi parliamo di intelligenza artificiale."),
    ]
    result = TranscriptionResult(
        segments=segments,
        full_text="Buongiorno a tutti. Oggi parliamo di intelligenza artificiale.",
        detected_language="it",
        language_probability=0.99,
        audio_duration=300.0,
        elapsed_time=15.0,
    )

    output = export_transcript_to_docx(
        output_path=docx_file,
        track_info=track_info,
        slice_config=slice_cfg,
        transcription_config=tx_cfg,
        result=result,
    )

    assert output.exists()
    assert output == docx_file

    # Load back with python-docx and inspect document contents
    doc = Document(str(docx_file))

    # Verify tables (metadata table)
    assert len(doc.tables) == 1
    meta_table = doc.tables[0]
    row_texts = [f"{row.cells[0].text}: {row.cells[1].text}" for row in meta_table.rows]
    joined_meta = " | ".join(row_texts)
    assert "lecture.mp3" in joined_meta
    assert "large-v3-turbo" in joined_meta
    assert "it" in joined_meta

    # Verify paragraphs contain the transcript and timestamps
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Audio Transcription" in doc_text
    assert "Transcript" in doc_text
    assert "Buongiorno a tutti." in doc_text
    assert "Oggi parliamo di intelligenza artificiale." in doc_text
    # Timestamps should be offset by slice start (60s -> 01:00)
    assert "[01:00 → 01:04]" in doc_text


def test_export_transcript_to_docx_without_timestamps(tmp_path: Path):
    docx_file = tmp_path / "transcript_no_ts.docx"

    track_info = AudioInfo(
        path=Path("/music/talk.wav"),
        filename="talk.wav",
        duration=60.0,
    )
    slice_cfg = SliceConfig(mode=SliceMode.FULL, start_sec=0.0, end_sec=60.0)
    tx_cfg = TranscriptionConfig(
        model_name="large-v3-turbo",
        include_timestamps_in_docx=False,
    )

    segments = [
        TranscriptSegment(start=0.0, end=2.0, text="First sentence."),
        TranscriptSegment(start=2.5, end=5.0, text="Second sentence."),
    ]
    result = TranscriptionResult(
        segments=segments,
        full_text="First sentence. Second sentence.",
        detected_language="en",
        language_probability=0.95,
        audio_duration=60.0,
        elapsed_time=5.0,
    )

    export_transcript_to_docx(
        output_path=docx_file,
        track_info=track_info,
        slice_config=slice_cfg,
        transcription_config=tx_cfg,
        result=result,
    )

    assert docx_file.exists()
    doc = Document(str(docx_file))
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    assert "First sentence." in doc_text
    assert "Second sentence." in doc_text
    assert "[" not in doc_text  # No timestamp brackets in body paragraphs


def test_export_transcript_to_docx_polished(tmp_path: Path):
    docx_file = tmp_path / "transcript_polished.docx"

    track_info = AudioInfo(
        path=Path("/recordings/meeting.m4a"),
        filename="meeting.m4a",
        duration=120.0,
    )
    slice_cfg = SliceConfig(mode=SliceMode.FULL, start_sec=0.0, end_sec=120.0)
    tx_cfg = TranscriptionConfig(
        model_name="large-v3-turbo",
        include_timestamps_in_docx=True,
    )

    segments = [
        TranscriptSegment(start=0.0, end=4.0, text="Ehm, so today we talk about the architecture."),
        TranscriptSegment(start=4.5, end=8.0, text="And, you know, we decided to use fast components."),
    ]
    result = TranscriptionResult(
        segments=segments,
        full_text="Ehm, so today we talk about the architecture. And, you know, we decided to use fast components.",
        detected_language="en",
        language_probability=0.98,
        audio_duration=120.0,
        elapsed_time=8.0,
        polished_text="Today we discussed the system architecture. We selected high-performance components.",
        llm_model_used="deepseek-v4.1-flash",
        llm_instructions_name="Default Instructions",
    )

    export_transcript_to_docx(
        output_path=docx_file,
        track_info=track_info,
        slice_config=slice_cfg,
        transcription_config=tx_cfg,
        result=result,
    )

    assert docx_file.exists()
    doc = Document(str(docx_file))

    # Check metadata table contains LLM details
    meta_table = doc.tables[0]
    meta_text = " | ".join(f"{r.cells[0].text}: {r.cells[1].text}" for r in meta_table.rows)
    assert "deepseek-v4.1-flash" in meta_text
    assert "Default Instructions" in meta_text

    # Check document paragraphs contain polished text and appendix
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Polished & Refined Document" in doc_text
    assert "Today we discussed the system architecture." in doc_text
    assert "Appendix: Original Raw Transcript (with Timestamps)" in doc_text
    assert "[00:00 → 00:04]" in doc_text
